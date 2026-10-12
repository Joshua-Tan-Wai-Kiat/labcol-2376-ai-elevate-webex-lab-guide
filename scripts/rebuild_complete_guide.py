"""Rebuild the public MkDocs guide from the latest LABCOL-2376 Word source.

The Word document is treated as source content only. This script extracts its
instructional text and embedded figures; it does not execute any lab command.
"""
from __future__ import annotations

import re
import shutil
from html import escape
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\joshutan\Downloads\CLMEL26 - LABCOL-2376 - Next-Gen Calling--AI-Elevate with Webex - Lab Guide (1).docx")
DOCS = ROOT / "docs"
ASSETS = DOCS / "assets" / "source-guide"

PAGE_RANGES = {
    "index.md": (19, 30),
    "01-access-the-lab.md": (31, 77),
    "02-webex-calling-org.md": (77, 167),
    "03-local-gateway.md": (167, 253),
    "04-captions-transcripts.md": (253, 327),
    "05-recording-summaries.md": (327, 372),
}

BULLET_PARAGRAPHS = {25, 26, 260, 261, 262, 363, 364, 367}
SKIP_PARAGRAPHS = {29}
OMIT_IMAGES = {
    "image4.png",   # old LABCOL-2416 session ID and VPN credentials
    "image5.png",   # example VPN username
    "image6.png",   # old LABCOL-2416 topology credentials
    "image8.png",   # Webex password and admin address
    "image9.png",   # old LABCOL-2416 session ID and DNS names
    "image10.png",  # credential-file icon adds no instructional detail
    "image11.png",  # Webex password and admin address
    "image17.png",  # bulk list of example phone numbers
    "image18.png",  # location street address
    "image19.png",  # location identifier
    "image20.png",  # phone number and organization domain
    "image21.png",  # phone number and location identifier
    "image22.png",  # identifiable profile photo and meeting domain
    "image23.png",  # user email and account URL identifier
    "image24.png",  # user email and phone numbers
    "image25.png",  # identifiable profile photo and meeting domain
    "image29.png",  # duplicated by the reviewed Local Gateway figure
    "image30.png",  # duplicated by the reviewed Local Gateway figure
    "image31.png",  # duplicated by the reviewed Local Gateway figure
    "image32.png",  # registrar domain, trunk identifier, and profile photo
    "image33.png",  # organization identifier
    "image36.png",  # phone number and location identifier
    "image38.png",  # identifiable profile photo
    "image40.png",  # user email, phone number, and profile photo
    "image41.png",  # identifiable profile photo
    "image42.png",  # organization domain and profile photo
    "image43.png",  # identifiable profile photo
    "image45.png",  # identifiable profile photo and call number
    "image46.png",  # identifiable profile photo and phone number
    "image47.png",  # dialable call number
    "image48.png",  # dialable call number
    "image49.png",  # dialable call number
    "image50.png",  # identifiable profile photo
    "image51.png",  # user email, organization domain, and profile photo
    "image52.png",  # pod-specific DID mapping
    "image53.jpeg", # decorative inline icon
    "image54.png",  # user phone number and profile photo
    "image55.jpeg", # decorative inline icon
    "image56.png",  # user phone number and profile photo
    "image57.jpeg", # decorative inline icon
    "image58.png",  # recorded call number
    "image59.png",  # call number and identifiable profile photo
}
SAFE_TEXT_REPLACEMENTS = {
    33: "Record the Host, Username, Password, and Session ID displayed for your assigned session.",
    34: "AnyConnect (Cisco Secure Client) Host: use the host shown in your session details.",
    35: "AnyConnect (Cisco Secure Client) Username: use the username shown in your session details.",
    36: "AnyConnect (Cisco Secure Client) Password: use the password shown in your session details.",
    64: "DNS Address / FQDN of your Webex organization: use the top-level FQDN shown for your assigned pod.",
    66: "Session ID: use the ID shown in your session details.",
    68: "Control Hub Username: cholland@&lt;FQDN&gt;, using your assigned top-level FQDN.",
    70: "Control Hub Password: dCloud + the last 4 digits of your Session ID + !",
}


def clean(text: str) -> str:
    text = text.replace("\ufffd", "'")
    text = re.sub(r"\s+", " ", text).strip()
    return text.replace("<FQDN>", "&lt;FQDN&gt;")


def image_names(document: Document, paragraph) -> list[str]:
    names: list[str] = []
    for blip in paragraph._p.xpath(".//a:blip"):
        rel_id = blip.get(qn("r:embed"))
        if rel_id and rel_id in document.part.rels:
            name = Path(str(document.part.rels[rel_id].target_ref)).name
            if name == "image26.tmp":
                name = "image26.png"
            names.append(name)
    return names


def extract_images(document: Document) -> None:
    if ASSETS.exists():
        shutil.rmtree(ASSETS)
    ASSETS.mkdir(parents=True, exist_ok=True)
    with ZipFile(SOURCE) as archive:
        for paragraph in document.paragraphs:
            for blip in paragraph._p.xpath(".//a:blip"):
                rel_id = blip.get(qn("r:embed"))
                if not rel_id or rel_id not in document.part.rels:
                    continue
                source_name = Path(str(document.part.rels[rel_id].target_ref)).name
                output_name = "image26.png" if source_name == "image26.tmp" else source_name
                if output_name in OMIT_IMAGES:
                    continue
                member = f"word/media/{source_name}"
                with archive.open(member) as source, (ASSETS / output_name).open("wb") as target:
                    shutil.copyfileobj(source, target)


def figure(name: str, context: str) -> str:
    alt = clean(context) or "Lab guide screenshot"
    alt = escape(alt[:180], quote=True)
    number = re.search(r"\d+", name)
    label = f"Source guide screenshot {number.group(0)}" if number else "Source guide screenshot"
    return (
        '<figure class="lab-figure">\n'
        f'  <img src="../assets/source-guide/{name}" alt="{alt}">\n'
        f"  <figcaption>{label}. Use the values assigned to your own lab pod.</figcaption>\n"
        "</figure>\n"
    )


PUTTY_BLOCK = """### Generate the Local Gateway configuration

1. Continuing on **Workstation 1**, minimize all applications and find the PowerShell script named `CL_Add_Trunk.ps1` on the desktop.
2. Right-click the file and select **Run with PowerShell**.

<figure class="lab-figure">
  <img src="../assets/local-gateway/01-run-trunk-script.png" alt="Right-click CL_Add_Trunk.ps1 and select Run with PowerShell">
  <figcaption>Run the trunk-creation script from the Workstation 1 desktop.</figcaption>
</figure>

3. Wait while the script executes. The PowerShell window closes automatically when it finishes.
4. Confirm that `CL_LGW_Config-Ready.txt` appears on the desktop, then double-click it to open the generated CUBE configuration.

### Connect to the LGW with PuTTY

1. Open **PuTTY** from the Workstation 1 taskbar or desktop.
2. In **Host Name (or IP address)**, enter `198.18.133.226`.
3. Leave **Port** set to `22`, select **SSH**, and click **Open**.
4. If PuTTY displays a host-key warning, confirm it only after checking that the address matches your assigned lab pod.
5. At the login prompts, enter username `admin` and password `dCloud123!`.

    PuTTY does not display password characters while you type. Press **Enter** after the password and wait for the `LGW#` prompt.

!!! warning "Source screenshot address mismatch"
    The written source step specifies `198.18.133.226`, but the supplied configuration screenshot shows `198.18.133.227` in its PuTTY title bar. Use the address assigned to your lab pod. Do not infer the correct address from the screenshot.

### Paste and apply the generated configuration

1. In `CL_LGW_Config-Ready.txt`, press **Ctrl+A**, then **Ctrl+C** to copy the entire configuration.
2. Return to the authenticated PuTTY window and right-click once inside the terminal to paste. The commands begin running immediately.
3. Wait until the `LGW#` prompt returns. Scroll through the terminal output and check for command errors.
4. You may ignore only the final save error shown after `copy run start`; do not ignore errors from the configuration commands.

<figure class="lab-figure">
  <img src="../assets/local-gateway/02-paste-lgw-config.png" alt="Generated LGW configuration selected in Notepad and pasted into an authenticated PuTTY session">
  <figcaption>Copy the complete generated configuration and paste it into the authenticated LGW session. This supplied screenshot shows a sample pod address; use your assigned address.</figcaption>
</figure>

### Verify Webex Calling registration

1. Wait **2 to 3 minutes** for the LGW to initiate registration with Webex Calling.
2. At the `LGW#` prompt, run `show sip-ua register status`.
3. In the command output, confirm that the **reg** column reports `yes`. If it does not, recheck the pasted Voice Class Tenant and SIP profile configuration before continuing.

<figure class="lab-figure">
  <img src="../assets/local-gateway/03-verify-registration.png" alt="Output of show sip-ua register status with the registration status highlighted as yes">
  <figcaption>The LGW is registered when the <strong>reg</strong> column shows <code>yes</code>.</figcaption>
</figure>

Once the registration status is `yes`, continue to Module 2b.
"""


def render_page(document: Document, filename: str, start: int, end: int) -> str:
    lines: list[str] = []
    last_context = "Lab guide screenshot"
    first_heading = True
    inserted_putty = False
    ordered_counter = 0

    if filename != "index.md":
        lines.extend([
            "> **Lab data notice:** Values in screenshots are examples from the source guide. Use the credentials, addresses, phone numbers, and session details assigned to your own pod.",
            "",
        ])

    for index in range(start, min(end, len(document.paragraphs))):
        paragraph = document.paragraphs[index]
        text = SAFE_TEXT_REPLACEMENTS.get(index, clean(paragraph.text))
        images = [name for name in image_names(document, paragraph) if name not in OMIT_IMAGES]

        if filename == "03-local-gateway.md" and 185 <= index <= 209:
            if not inserted_putty:
                lines.extend([PUTTY_BLOCK.rstrip(), ""])
                inserted_putty = True
            continue
        if index in SKIP_PARAGRAPHS:
            continue

        if text:
            last_context = text
            style = paragraph.style.name
            if style == "Heading 1":
                ordered_counter = 0
                level = "#" if first_heading else "##"
                lines.extend([f"{level} {text}", ""])
                first_heading = False
            elif style == "Heading 2":
                ordered_counter = 0
                lines.extend([f"## {text}", ""])
            elif style == "Heading 3":
                ordered_counter = 0
                lines.extend([f"### {text}", ""])
            elif style == "List Paragraph":
                if index in BULLET_PARAGRAPHS:
                    marker = "-"
                    ordered_counter = 0
                else:
                    ordered_counter += 1
                    marker = f"{ordered_counter}."
                lines.append(f"{marker} {text}")
            else:
                lines.extend([text, ""])

        for name in images:
            if name in {"image29.png", "image30.png", "image31.png"} and filename == "03-local-gateway.md":
                continue
            lines.extend([figure(name, last_context).rstrip(), ""])

    title_prefix = "" if filename == "index.md" else ""
    return title_prefix + "\n".join(lines).strip() + "\n"


def add_index_front_matter(body: str) -> str:
    header = """# Next Gen Calling AI Elevate with Webex

<p class="lab-meta">Cisco Live Melbourne 2026 | LABCOL-2376</p>

This online edition contains the complete hands-on workflow and instructional screenshots from the current lab guide.

**Authors:** Hussain Ali and Omer Ilyas, Technical Marketing Engineers

## Lab modules

- [Access your lab](01-access-the-lab.md)
- [Set up the Webex Calling organization](02-webex-calling-org.md)
- [Configure the Local Gateway](03-local-gateway.md)
- [Test captions and call transcripts](04-captions-transcripts.md)
- [Test call recording summaries](05-recording-summaries.md)

## Webex Calling AI Overview

"""
    body = re.sub(r"^# Webex Calling AI Overview\s*", "", body)
    return header + body


def main() -> None:
    document = Document(SOURCE)
    extract_images(document)
    for filename, (start, end) in PAGE_RANGES.items():
        content = render_page(document, filename, start, end)
        if filename == "index.md":
            content = add_index_front_matter(content)
        (DOCS / filename).write_text(content, encoding="utf-8")
        print(f"Wrote {filename}")


if __name__ == "__main__":
    main()
